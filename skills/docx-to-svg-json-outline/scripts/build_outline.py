#!/usr/bin/env python3
"""Generate PPT SVG JSON outlines and image placeholders from a chapter DOCX."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from PIL import Image
except ImportError as exc:  # pragma: no cover - environment guidance
    raise SystemExit(
        "python-docx and Pillow are required. Use the bundled Codex workspace Python runtime."
    ) from exc


OPEN_MARKER = "[PPT_OUTLINE]"
CLOSE_MARKER = "[/PPT_OUTLINE]"
NUMBERED_HEADING = re.compile(r"^\s*(\d+)\.(\d+)\s*")
DIRECT_SVG_TYPES = {".svg"}
DIRECT_RASTER_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


@dataclass
class ImageAsset:
    blob: bytes
    extension: str
    ratio: float
    source_name: str
    filename: str = ""

    @property
    def placeholder(self) -> str:
        return f"图片占位：{self.filename}（宽高比 {self.ratio:.2f}:1）"


@dataclass
class Page:
    title: str
    images: list[ImageAsset] = field(default_factory=list)


@dataclass
class Section:
    title: str
    pages: list[Page] = field(default_factory=list)
    cover_images: list[ImageAsset] = field(default_factory=list)
    chapter: int | None = None
    section: int | None = None
    filename: str = ""

    @property
    def images(self) -> list[ImageAsset]:
        return self.cover_images + [image for page in self.pages for image in page.images]


def normalize_style(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip()).casefold()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-docx", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--cover-title", required=True)
    parser.add_argument("--toc-title", default="目录")
    parser.add_argument("--chapter-number", type=int)
    parser.add_argument("--heading1-style", action="append", default=[])
    parser.add_argument("--heading2-style", action="append", default=[])
    parser.add_argument("--plan", action="store_true", help="Print mapping without writing files")
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def displayed_ratio(element) -> float | None:
    drawing = element.xpath("ancestor::w:drawing[1]")
    if not drawing:
        return None
    extents = drawing[0].xpath(".//wp:extent")
    if not extents:
        return None
    width = int(extents[0].get("cx", "0"))
    height = int(extents[0].get("cy", "0"))
    return width / height if width > 0 and height > 0 else None


def raster_details(blob: bytes) -> tuple[float, Image.Image]:
    image = Image.open(BytesIO(blob))
    image.seek(0)
    width, height = image.size
    if width <= 0 or height <= 0:
        raise ValueError("Image has invalid dimensions")
    return width / height, image.copy()


def normalize_image(blob: bytes, suffix: str) -> tuple[bytes, str, float]:
    suffix = suffix.casefold()
    if suffix in DIRECT_SVG_TYPES:
        return blob, suffix, 0.0

    intrinsic_ratio, image = raster_details(blob)
    if suffix in DIRECT_RASTER_TYPES:
        return blob, ".jpg" if suffix == ".jpeg" else suffix, intrinsic_ratio

    output = BytesIO()
    if image.mode not in {"RGB", "RGBA", "L", "LA", "P"}:
        image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
    image.save(output, format="PNG")
    return output.getvalue(), ".png", intrinsic_ratio


def related_image(paragraph, relationship_id: str, ratio: float | None) -> ImageAsset:
    part = paragraph.part.related_parts.get(relationship_id)
    if part is None:
        raise ValueError(f"Linked or missing image relationship: {relationship_id}")
    source_name = Path(str(part.partname)).name
    blob, extension, intrinsic_ratio = normalize_image(part.blob, Path(source_name).suffix)
    final_ratio = ratio or intrinsic_ratio
    if final_ratio <= 0:
        raise ValueError(f"Cannot determine image ratio: {source_name}")
    return ImageAsset(blob, extension, final_ratio, source_name)


def paragraph_images(paragraph) -> list[ImageAsset]:
    images: list[ImageAsset] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relationship_id = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
        if relationship_id:
            images.append(related_image(paragraph, relationship_id, displayed_ratio(blip)))
    for image_data in paragraph._p.xpath(".//*[local-name()='imagedata']"):
        relationship_id = image_data.get(qn("r:id"))
        if relationship_id:
            images.append(related_image(paragraph, relationship_id, None))
    return images


def extract_sections(
    input_docx: Path, heading1_styles: set[str], heading2_styles: set[str]
) -> tuple[list[Section], list[str], list[ImageAsset]]:
    document = Document(input_docx)
    sections: list[Section] = []
    body_text: list[str] = []
    directory_images: list[ImageAsset] = []
    current_section: Section | None = None
    current_page: Page | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = normalize_style(paragraph.style.name or "")

        if text and style in heading1_styles:
            match = NUMBERED_HEADING.match(text)
            current_section = Section(
                title=text,
                chapter=int(match.group(1)) if match else None,
                section=int(match.group(2)) if match else None,
            )
            sections.append(current_section)
            current_page = None
        elif text and style in heading2_styles:
            if current_section is not None:
                current_page = Page(text)
                current_section.pages.append(current_page)
        elif text:
            body_text.append(text)

        images = paragraph_images(paragraph)
        if not images:
            continue
        if current_section is None:
            directory_images.extend(images)
        elif current_page is None:
            current_section.cover_images.extend(images)
        else:
            current_page.images.extend(images)

    if not sections:
        raise ValueError("No nonempty Heading 1 paragraphs were found")
    return sections, body_text, directory_images


def resolve_chapter_number(sections: list[Section], explicit: int | None) -> int:
    numbered_chapters = [item.chapter for item in sections if item.chapter is not None]
    if explicit is not None:
        if numbered_chapters and any(value != explicit for value in numbered_chapters):
            raise ValueError("Numbered headings conflict with --chapter-number")
        return explicit
    if not numbered_chapters:
        raise ValueError("Cannot infer chapter number; pass --chapter-number")
    if len(set(numbered_chapters)) != 1:
        raise ValueError("DOCX contains numbered headings from multiple chapters")
    return numbered_chapters[0]


def assign_filenames(sections: list[Section], chapter_number: int) -> None:
    base = chapter_number * 100
    current_section: int | None = None
    occurrences: defaultdict[int | None, int] = defaultdict(int)

    for item in sections:
        if item.chapter is not None:
            if item.chapter != chapter_number:
                raise ValueError(f"Heading belongs to chapter {item.chapter}: {item.title}")
            current_section = item.section
        key = current_section
        occurrences[key] += 1
        if key is None:
            item.filename = f"{base}-{occurrences[key] + 1}.json"
        else:
            stem = base + key
            occurrence = occurrences[key]
            item.filename = f"{stem}.json" if occurrence == 1 else f"{stem}-{occurrence}.json"

        json_stem = Path(item.filename).stem
        for index, image in enumerate(item.images, start=1):
            image.filename = f"{json_stem}-image-{index:02d}{image.extension}"


def assign_directory_image_filenames(images: list[ImageAsset], base: int) -> None:
    for index, image in enumerate(images, start=1):
        image.filename = f"{base}-image-{index:02d}{image.extension}"


def directory_entries(sections: list[Section], chapter_number: int) -> list[str]:
    first_by_section: dict[int, str] = {}
    for item in sections:
        if item.chapter == chapter_number and item.section is not None:
            first_by_section.setdefault(item.section, item.title)
    return [first_by_section[key] for key in sorted(first_by_section)]


def placeholders(images: list[ImageAsset]) -> list[str]:
    return [image.placeholder for image in images]


def content_outline(section: Section) -> dict:
    outline: dict = {
        "cover": {"title": section.title, "content": placeholders(section.cover_images)},
    }
    if section.pages:
        outline["parts"] = [
            {
                "pages": [
                    {"title": page.title, "content": placeholders(page.images)}
                    for page in section.pages
                ]
            }
        ]
    return {"ppt_outline": outline}


def directory_outline(
    cover_title: str, toc_title: str, entries: list[str], images: list[ImageAsset]
) -> dict:
    return {
        "ppt_outline": {
            "cover": {"title": cover_title, "content": placeholders(images)},
            "table_of_contents": {"title": toc_title, "content": entries},
        }
    }


def serialize(payload: dict) -> str:
    body = json.dumps(payload, ensure_ascii=False, indent=2)
    return f"{OPEN_MARKER}\n{body}\n{CLOSE_MARKER}\n"


def parse_generated(path: Path) -> dict:
    text = path.read_text(encoding="utf-8").strip()
    if not (text.startswith(OPEN_MARKER) and text.endswith(CLOSE_MARKER)):
        raise ValueError(f"Markers missing in {path}")
    body = text[len(OPEN_MARKER) : -len(CLOSE_MARKER)].strip()
    return json.loads(body)


def all_images(sections: list[Section], directory_images: list[ImageAsset]) -> list[ImageAsset]:
    return directory_images + [image for section in sections for image in section.images]


def validate_output(
    output_dir: Path,
    directory_filename: str,
    sections: list[Section],
    expected_entries: list[str],
    body_text: list[str],
    directory_images: list[ImageAsset],
) -> None:
    content_paths = [output_dir / item.filename for item in sections]
    payloads = [parse_generated(path) for path in content_paths]
    directory = parse_generated(output_dir / directory_filename)
    covers = [payload["ppt_outline"]["cover"]["title"] for payload in payloads]
    pages: list[str] = []

    for payload in payloads:
        outline = payload["ppt_outline"]
        for content in [outline["cover"].get("content", [])] + [
            page.get("content", [])
            for part in outline.get("parts", [])
            for page in part.get("pages", [])
        ]:
            if any(not value.startswith("图片占位：") for value in content):
                raise ValueError("Non-image body content found in content array")
        pages.extend(
            page["title"]
            for part in outline.get("parts", [])
            for page in part.get("pages", [])
        )

    if Counter(covers) != Counter(item.title for item in sections):
        raise ValueError("Heading 1 coverage check failed")
    if Counter(pages) != Counter(page.title for item in sections for page in item.pages):
        raise ValueError("Heading 2 coverage check failed")
    actual_entries = directory["ppt_outline"]["table_of_contents"]["content"]
    if actual_entries != expected_entries:
        raise ValueError("Directory entry check failed")

    expected_images = all_images(sections, directory_images)
    for image in expected_images:
        path = output_dir / image.filename
        if not path.is_file() or path.stat().st_size == 0:
            raise ValueError(f"Extracted image is missing or empty: {image.filename}")
    generated_json = "\n".join(path.read_text(encoding="utf-8") for path in content_paths)
    generated_json += (output_dir / directory_filename).read_text(encoding="utf-8")
    for image in expected_images:
        if image.placeholder not in generated_json:
            raise ValueError(f"Image placeholder is missing: {image.filename}")

    leaked = [text for text in body_text if len(text) >= 12 and text in generated_json]
    if leaked:
        raise ValueError(f"Body content leaked into output: {leaked[0][:80]}")


def main() -> int:
    args = parse_args()
    if not args.input_docx.is_file():
        raise SystemExit(f"Input DOCX not found: {args.input_docx}")

    heading1_styles = {normalize_style("Heading 1"), normalize_style("标题 1")}
    heading2_styles = {normalize_style("Heading 2"), normalize_style("标题 2")}
    heading1_styles.update(normalize_style(value) for value in args.heading1_style)
    heading2_styles.update(normalize_style(value) for value in args.heading2_style)

    try:
        sections, body_text, directory_images = extract_sections(
            args.input_docx, heading1_styles, heading2_styles
        )
        chapter_number = resolve_chapter_number(sections, args.chapter_number)
        assign_filenames(sections, chapter_number)
        base = chapter_number * 100
        assign_directory_image_filenames(directory_images, base)
        entries = directory_entries(sections, chapter_number)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc

    directory_filename = f"{base}.json"
    print(
        f"directory: {directory_filename} | {args.cover_title} | "
        f"entries={len(entries)} images={len(directory_images)}"
    )
    for item in sections:
        print(
            f"{item.filename}: {item.title} | pages={len(item.pages)} "
            f"images={len(item.images)}"
        )
        for image in item.images:
            print(f"  {image.filename}: ratio={image.ratio:.2f}:1 source={image.source_name}")
    image_count = len(all_images(sections, directory_images))
    print(
        f"summary: files={len(sections) + 1} heading1={len(sections)} "
        f"heading2={sum(len(item.pages) for item in sections)} images={image_count}"
    )
    if args.plan:
        return 0

    json_paths = [args.output_dir / directory_filename] + [
        args.output_dir / item.filename for item in sections
    ]
    image_paths = [args.output_dir / image.filename for image in all_images(sections, directory_images)]
    collisions = [path for path in json_paths + image_paths if path.exists()]
    if collisions and not args.overwrite:
        names = ", ".join(path.name for path in collisions[:5])
        raise SystemExit(f"Refusing to overwrite existing files: {names}")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / directory_filename).write_text(
        serialize(directory_outline(args.cover_title, args.toc_title, entries, directory_images)),
        encoding="utf-8",
    )
    for item in sections:
        (args.output_dir / item.filename).write_text(
            serialize(content_outline(item)), encoding="utf-8"
        )
    for image in all_images(sections, directory_images):
        (args.output_dir / image.filename).write_bytes(image.blob)

    try:
        validate_output(
            args.output_dir,
            directory_filename,
            sections,
            entries,
            body_text,
            directory_images,
        )
    except ValueError as exc:
        raise SystemExit(f"Validation failed: {exc}") from exc

    print("validation: PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
